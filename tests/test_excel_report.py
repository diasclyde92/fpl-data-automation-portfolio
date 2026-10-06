"""Unit tests for the ExcelReport generation layer using openpyxl."""

from pathlib import Path

import openpyxl
import pytest

from src.analytics.book_analytics import BookAnalytics
from src.models.book import Book
from src.reporting.excel_report import ExcelReport
from src.storage.sqlite_storage import SQLiteStorage


@pytest.fixture
def empty_storage(tmp_path: Path) -> SQLiteStorage:
    """Fixture with an empty database."""
    return SQLiteStorage(db_path=tmp_path / "empty.db")


@pytest.fixture
def populated_storage(tmp_path: Path) -> SQLiteStorage:
    """Fixture with a populated database."""
    storage = SQLiteStorage(db_path=tmp_path / "test_populated.db")
    books = [
        Book(
            title="Book One",
            price=15.00,
            rating=2,
            availability="In stock",
            detail_url="http://books.toscrape.com/b1",
        ),
        Book(
            title="Book Two",
            price=25.00,
            rating=4,
            availability="In stock",
            detail_url="http://books.toscrape.com/b2",
        ),
        Book(
            title="Book Three",
            price=35.00,
            rating=5,
            availability="Out of stock",
            detail_url="http://books.toscrape.com/b3",
        ),
    ]
    storage.insert_books(books)
    return storage


def test_excel_report_generation_and_sheet_names(populated_storage: SQLiteStorage, tmp_path: Path):
    """Verify workbook generates all expected sheets with correct titles."""
    analytics = BookAnalytics(storage=populated_storage)
    report = ExcelReport(analytics=analytics)
    output_path = tmp_path / "test_report.xlsx"

    result_path = report.generate(output_path=output_path)

    assert result_path.exists()
    wb = openpyxl.load_workbook(str(result_path))

    expected_sheets = ["Summary", "Books Data", "Price Analysis", "Rating Analysis", "Availability"]
    assert wb.sheetnames == expected_sheets


def test_excel_report_summary_kpis(populated_storage: SQLiteStorage, tmp_path: Path):
    """Verify that the Summary sheet contains calculated KPIs matching data."""
    analytics = BookAnalytics(storage=populated_storage)
    report = ExcelReport(analytics=analytics)
    output_path = tmp_path / "kpi_report.xlsx"

    report.generate(output_path=output_path)
    wb = openpyxl.load_workbook(str(output_path), data_only=True)
    ws = wb["Summary"]

    # In our template:
    # Row 8: Total Catalog Items = 3
    # Row 9: Average Price = 25.0
    # Row 10: Median Price = 25.0
    # Row 11: Minimum Price = 15.0
    # Row 12: Maximum Price = 35.0
    # Row 13: Average Rating = 3.67
    assert ws["B8"].value == "Total Catalog Items"
    assert ws["C8"].value == 3
    assert ws["B9"].value == "Average Price (GBP)"
    assert ws["C9"].value == 25.0
    assert ws["B10"].value == "Median Price (GBP)"
    assert ws["C10"].value == 25.0
    assert ws["B11"].value == "Minimum Price (GBP)"
    assert ws["C11"].value == 15.0
    assert ws["B12"].value == "Maximum Price (GBP)"
    assert ws["C12"].value == 35.0


def test_excel_report_books_data_rows(populated_storage: SQLiteStorage, tmp_path: Path):
    """Verify Books Data table contains header row plus exactly 3 data rows."""
    analytics = BookAnalytics(storage=populated_storage)
    report = ExcelReport(analytics=analytics)
    output_path = tmp_path / "data_report.xlsx"

    report.generate(output_path=output_path)
    wb = openpyxl.load_workbook(str(output_path))
    ws = wb["Books Data"]

    # Header in row 1
    assert ws["A1"].value == "ID"
    assert ws["B1"].value == "Title"
    assert ws["C1"].value == "Price (GBP)"

    # Rows 2 to 4 must contain our books
    assert ws.max_row == 4
    assert ws["B2"].value == "Book One"
    assert ws["C2"].value == 15.0
    assert ws["D2"].value == 2
    assert ws["E2"].value == "In stock"
    assert ws["B4"].value == "Book Three"
    assert ws["C4"].value == 35.0


def test_excel_report_price_analysis(populated_storage: SQLiteStorage, tmp_path: Path):
    """Verify Price Analysis sheet metrics and top items."""
    analytics = BookAnalytics(storage=populated_storage)
    report = ExcelReport(analytics=analytics)
    output_path = tmp_path / "price_report.xlsx"

    report.generate(output_path=output_path)
    wb = openpyxl.load_workbook(str(output_path))
    ws = wb["Price Analysis"]

    # Metrics table
    assert ws["B8"].value == "Average Book Price"
    assert ws["C8"].value == 25.0
    assert ws["C11"].value == 35.0  # Max price

    # Top expensive items start at row 17
    assert ws["B17"].value == "Book Three"
    assert ws["C17"].value == 35.0


def test_excel_report_rating_analysis(populated_storage: SQLiteStorage, tmp_path: Path):
    """Verify Rating Analysis distribution table."""
    analytics = BookAnalytics(storage=populated_storage)
    report = ExcelReport(analytics=analytics)
    output_path = tmp_path / "rating_report.xlsx"

    report.generate(output_path=output_path)
    wb = openpyxl.load_workbook(str(output_path))
    ws = wb["Rating Analysis"]

    # Rows 8-12 correspond to 1 to 5 stars
    # 2-star: 1 book, 4-star: 1 book, 5-star: 1 book
    assert ws["B9"].value == "2 Star"
    assert ws["C9"].value == 1
    assert ws["B11"].value == "4 Star"
    assert ws["C11"].value == 1
    assert ws["B12"].value == "5 Star"
    assert ws["C12"].value == 1
    assert ws["B8"].value == "1 Star"
    assert ws["C8"].value == 0


def test_excel_report_availability_sheet(populated_storage: SQLiteStorage, tmp_path: Path):
    """Verify Availability breakdown sheet."""
    analytics = BookAnalytics(storage=populated_storage)
    report = ExcelReport(analytics=analytics)
    output_path = tmp_path / "avail_report.xlsx"

    report.generate(output_path=output_path)
    wb = openpyxl.load_workbook(str(output_path))
    ws = wb["Availability"]

    # Table starts at row 8
    # "In stock": 2, "Out of stock": 1
    status_values = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=3).value for r in [8, 9]}
    assert status_values.get("In stock") == 2
    assert status_values.get("Out of stock") == 1


def test_excel_report_empty_database_handling(empty_storage: SQLiteStorage, tmp_path: Path):
    """Verify that generating report on empty database succeeds without crashing."""
    analytics = BookAnalytics(storage=empty_storage)
    report = ExcelReport(analytics=analytics)
    output_path = tmp_path / "empty_report.xlsx"

    result_path = report.generate(output_path=output_path)
    assert result_path.exists()

    wb = openpyxl.load_workbook(str(result_path))
    assert len(wb.sheetnames) == 5

    # Check Summary shows 0 items and EMPTY status
    ws_sum = wb["Summary"]
    assert ws_sum["C8"].value == 0
    assert ws_sum["F8"].value == "EMPTY"

    # Books Data has headers only
    ws_data = wb["Books Data"]
    assert ws_data.max_row == 1
