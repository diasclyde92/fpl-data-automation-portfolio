"""Excel report generator for books data analytics using openpyxl."""

import argparse
import logging
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.analytics.book_analytics import BookAnalytics
from src.storage.sqlite_storage import SQLiteStorage

logger = logging.getLogger(__name__)

# Professional corporate styling constants (Executive Navy Theme)
FONT_FAMILY = "Segoe UI"
COLOR_NAVY_DARK = "1B365D"
COLOR_NAVY_LIGHT = "E8EEF5"
COLOR_HEADER_FILL = "24426B"
COLOR_HEADER_TEXT = "FFFFFF"
COLOR_ZEBRA = "F9FBFC"
COLOR_BORDER = "D9D9D9"
COLOR_BORDER_STRONG = "1B365D"
COLOR_STATUS_HEALTHY = "2E7D32"
COLOR_STATUS_EMPTY = "C62828"

HEADER_FILL = PatternFill(start_color=COLOR_HEADER_FILL, end_color=COLOR_HEADER_FILL, fill_type="solid")
HEADER_FONT = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_HEADER_TEXT)
TITLE_FONT = Font(name=FONT_FAMILY, size=16, bold=True, color=COLOR_NAVY_DARK)
SUBTITLE_FONT = Font(name=FONT_FAMILY, size=10, italic=True, color="555555")
SECTION_FONT = Font(name=FONT_FAMILY, size=12, bold=True, color=COLOR_NAVY_DARK)
LABEL_FONT = Font(name=FONT_FAMILY, size=11, bold=True, color="333333")
DATA_FONT = Font(name=FONT_FAMILY, size=11, color="111111")
ZEBRA_FILL = PatternFill(start_color=COLOR_ZEBRA, end_color=COLOR_ZEBRA, fill_type="solid")
CARD_FILL = PatternFill(start_color=COLOR_NAVY_LIGHT, end_color=COLOR_NAVY_LIGHT, fill_type="solid")

THIN_SIDE = Side(border_style="thin", color=COLOR_BORDER)
THICK_BOTTOM = Side(border_style="medium", color=COLOR_BORDER_STRONG)
BORDER_ALL = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
BORDER_HEADER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THICK_BOTTOM)

CURRENCY_FORMAT = '£#,##0.00'
PERCENT_FORMAT = '0.0%'
INTEGER_FORMAT = '#,##0'
DATETIME_FORMAT = 'yyyy-mm-dd hh:mm:ss'


class ExcelReport:
    """Generates styled multi-tab Excel workbooks from BookAnalytics outputs."""

    def __init__(self, analytics: BookAnalytics) -> None:
        """Initialize ExcelReport with an analytics provider.

        Args:
            analytics: BookAnalytics instance.
        """
        self.analytics = analytics

    def generate(self, output_path: Path | str = "reports/exports/books_analytics_report.xlsx") -> Path:
        """Generate the complete multi-worksheet Excel workbook.

        Args:
            output_path: Path where the .xlsx file will be saved.

        Returns:
            Path object of the saved file.
        """
        target = Path(output_path)
        target.parent.mkdir(parents=True, exist_ok=True)

        logger.info("Generating Excel report from %s", self.analytics.storage.db_path)
        df = self.analytics.load_data()
        result = self.analytics.analyze()

        wb = openpyxl.Workbook()
        # Remove default empty sheet
        wb.remove(wb.active)

        self._build_summary_sheet(wb, result)
        self._build_books_data_sheet(wb, df)
        self._build_price_analysis_sheet(wb, df, result)
        self._build_rating_analysis_sheet(wb, df, result)
        self._build_availability_sheet(wb, df, result)

        wb.save(str(target))
        logger.info("Successfully generated and saved Excel report to %s", target)
        return target

    def _auto_fit_columns(self, ws, max_len_cap: int = 45, min_len: int = 12) -> None:
        """Auto-adjust worksheet column widths based on contents."""
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if "\n" in val_str:
                    val_str = max(val_str.split("\n"), key=len)
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max_len_cap, max(max_len + 3, min_len))

    # -------------------------------------------------------------------------
    # Sheet 1: Executive Summary
    # -------------------------------------------------------------------------
    def _build_summary_sheet(self, wb: openpyxl.Workbook, result) -> None:
        """Build executive summary worksheet with KPI cards, rankings, and quality audit."""
        ws = wb.create_sheet(title="Summary")
        ws.views.sheetView[0].showGridLines = True

        # Header Title
        ws["B2"] = "Books Catalog Executive Summary"
        ws["B2"].font = TITLE_FONT
        ws["B3"] = f"Automated Pipeline Deliverable | Source Database: {self.analytics.storage.db_path.name}"
        ws["B3"].font = SUBTITLE_FONT

        summary = result.summary
        quality = result.quality

        # KPI Block Table
        ws["B5"] = "Key Performance Indicators"
        ws["B5"].font = SECTION_FONT

        kpi_metrics = [
            ("Total Catalog Items", summary.total_books, INTEGER_FORMAT),
            ("Average Price (GBP)", summary.average_price, CURRENCY_FORMAT),
            ("Median Price (GBP)", summary.median_price, CURRENCY_FORMAT),
            ("Minimum Price (GBP)", summary.min_price, CURRENCY_FORMAT),
            ("Maximum Price (GBP)", summary.max_price, CURRENCY_FORMAT),
            ("Average Star Rating", summary.average_rating, "0.00"),
        ]

        # Table Header
        ws["B7"] = "Metric"
        ws["C7"] = "Value"
        ws["B7"].font = HEADER_FONT
        ws["C7"].font = HEADER_FONT
        ws["B7"].fill = HEADER_FILL
        ws["C7"].fill = HEADER_FILL
        ws["B7"].alignment = Alignment(horizontal="left")
        ws["C7"].alignment = Alignment(horizontal="right")
        ws["B7"].border = BORDER_HEADER
        ws["C7"].border = BORDER_HEADER

        row_idx = 8
        for label, val, num_fmt in kpi_metrics:
            cell_lbl = ws.cell(row=row_idx, column=2, value=label)
            cell_val = ws.cell(row=row_idx, column=3, value=val if val is not None else "N/A")

            cell_lbl.font = LABEL_FONT
            cell_val.font = DATA_FONT
            cell_lbl.border = BORDER_ALL
            cell_val.border = BORDER_ALL
            cell_val.alignment = Alignment(horizontal="right")

            if val is not None and isinstance(val, (int, float)):
                cell_val.number_format = num_fmt

            if row_idx % 2 == 1:
                cell_lbl.fill = ZEBRA_FILL
                cell_val.fill = ZEBRA_FILL

            row_idx += 1

        # Data Quality Audit Section
        ws["E5"] = "Data Quality & Health Audit"
        ws["E5"].font = SECTION_FONT

        ws["E7"] = "Quality Check"
        ws["F7"] = "Status / Count"
        ws["E7"].font = HEADER_FONT
        ws["F7"].font = HEADER_FONT
        ws["E7"].fill = HEADER_FILL
        ws["F7"].fill = HEADER_FILL
        ws["E7"].alignment = Alignment(horizontal="left")
        ws["F7"].alignment = Alignment(horizontal="right")
        ws["E7"].border = BORDER_HEADER
        ws["F7"].border = BORDER_HEADER

        status_text = "EMPTY" if quality.is_empty else "HEALTHY"
        status_color = COLOR_STATUS_EMPTY if quality.is_empty else COLOR_STATUS_HEALTHY

        quality_checks = [
            ("Dataset Status", status_text, None),
            ("Missing Prices", quality.missing_prices, INTEGER_FORMAT),
            ("Missing Ratings", quality.missing_ratings, INTEGER_FORMAT),
            ("Missing Stock Status", quality.missing_availability, INTEGER_FORMAT),
            ("Duplicate URLs", quality.duplicate_urls, INTEGER_FORMAT),
            ("Total Records Validated", quality.total_records, INTEGER_FORMAT),
        ]

        q_row = 8
        for label, val, num_fmt in quality_checks:
            c_lbl = ws.cell(row=q_row, column=5, value=label)
            c_val = ws.cell(row=q_row, column=6, value=val)

            c_lbl.font = LABEL_FONT
            c_val.font = Font(name=FONT_FAMILY, size=11, bold=(label == "Dataset Status"), color=status_color if label == "Dataset Status" else "111111")
            c_lbl.border = BORDER_ALL
            c_val.border = BORDER_ALL
            c_val.alignment = Alignment(horizontal="right")

            if num_fmt and isinstance(val, (int, float)):
                c_val.number_format = num_fmt

            if q_row % 2 == 1:
                c_lbl.fill = ZEBRA_FILL
                c_val.fill = ZEBRA_FILL

            q_row += 1

        # Executive Highlights: Top & Bottom Tables
        start_rank_row = 16
        ws.cell(row=start_rank_row, column=2, value="Top 5 Most Expensive Books").font = SECTION_FONT
        self._write_mini_rank_table(ws, start_row=start_rank_row + 2, start_col=2, items=result.top_expensive)

        ws.cell(row=start_rank_row, column=6, value="Top Rated Books").font = SECTION_FONT
        self._write_mini_rank_table(ws, start_row=start_rank_row + 2, start_col=6, items=result.top_rated)

        self._auto_fit_columns(ws)

    def _write_mini_rank_table(self, ws, start_row: int, start_col: int, items: list[dict]) -> None:
        """Write a compact ranked items table on summary sheet."""
        headers = ["Title", "Price", "Rating"]
        for c_idx, h in enumerate(headers):
            cell = ws.cell(row=start_row, column=start_col + c_idx, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = BORDER_HEADER
            cell.alignment = Alignment(horizontal="right" if h in ["Price", "Rating"] else "left")

        curr_row = start_row + 1
        if not items:
            cell = ws.cell(row=curr_row, column=start_col, value="(No data available)")
            cell.font = Font(name=FONT_FAMILY, italic=True)
            return

        for item in items:
            c1 = ws.cell(row=curr_row, column=start_col, value=item.get("title", "N/A"))
            c2 = ws.cell(row=curr_row, column=start_col + 1, value=item.get("price"))
            c3 = ws.cell(row=curr_row, column=start_col + 2, value=item.get("rating"))

            c1.font = DATA_FONT
            c2.font = DATA_FONT
            c3.font = DATA_FONT

            c1.border = BORDER_ALL
            c2.border = BORDER_ALL
            c3.border = BORDER_ALL

            c2.number_format = CURRENCY_FORMAT
            c2.alignment = Alignment(horizontal="right")
            c3.alignment = Alignment(horizontal="right")

            if curr_row % 2 == 1:
                c1.fill = ZEBRA_FILL
                c2.fill = ZEBRA_FILL
                c3.fill = ZEBRA_FILL

            curr_row += 1

    # -------------------------------------------------------------------------
    # Sheet 2: Books Data
    # -------------------------------------------------------------------------
    def _build_books_data_sheet(self, wb: openpyxl.Workbook, df) -> None:
        """Build raw/persisted data tabular sheet with autofilter and freeze panes."""
        ws = wb.create_sheet(title="Books Data")
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A2"

        headers = ["ID", "Title", "Price (GBP)", "Rating", "Availability", "Detail URL", "Scraped At (UTC)"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = BORDER_HEADER
            cell.alignment = Alignment(horizontal="center" if h in ["ID", "Rating"] else ("right" if "Price" in h else "left"))

        if not df.empty:
            for r_idx, row in df.iterrows():
                row_num = r_idx + 2
                c_id = ws.cell(row=row_num, column=1, value=int(row["id"]))
                c_title = ws.cell(row=row_num, column=2, value=str(row["title"]))
                c_price = ws.cell(row=row_num, column=3, value=float(row["price"]) if row["price"] is not None else None)
                c_rating = ws.cell(row=row_num, column=4, value=int(row["rating"]) if pd.notna(row["rating"]) else None)
                c_avail = ws.cell(row=row_num, column=5, value=str(row["availability"]))
                c_url = ws.cell(row=row_num, column=6, value=str(row["detail_url"]))
                c_scraped = ws.cell(row=row_num, column=7, value=str(row["scraped_at"]))

                for cell in (c_id, c_title, c_price, c_rating, c_avail, c_url, c_scraped):
                    cell.font = DATA_FONT
                    cell.border = BORDER_ALL
                    if row_num % 2 == 1:
                        cell.fill = ZEBRA_FILL

                c_id.alignment = Alignment(horizontal="center")
                c_price.number_format = CURRENCY_FORMAT
                c_price.alignment = Alignment(horizontal="right")
                c_rating.alignment = Alignment(horizontal="center")

                # Make URL clickable hyperlink
                if row["detail_url"]:
                    c_url.hyperlink = str(row["detail_url"])
                    c_url.font = Font(name=FONT_FAMILY, color="0000FF", underline="single")

            # Enable auto-filter
            last_col_letter = get_column_letter(len(headers))
            ws.auto_filter.ref = f"A1:{last_col_letter}{len(df) + 1}"

        self._auto_fit_columns(ws, max_len_cap=50)

    # -------------------------------------------------------------------------
    # Sheet 3: Price Analysis
    # -------------------------------------------------------------------------
    def _build_price_analysis_sheet(self, wb: openpyxl.Workbook, df, result) -> None:
        """Build price distribution and top/bottom price rankings with an Excel chart."""
        ws = wb.create_sheet(title="Price Analysis")
        ws.views.sheetView[0].showGridLines = True

        ws["B2"] = "Catalog Price Distribution & Rankings"
        ws["B2"].font = TITLE_FONT
        ws["B3"] = "Key valuation metrics and pricing extremes"
        ws["B3"].font = SUBTITLE_FONT

        # Price summary metrics
        ws["B5"] = "Valuation Metrics"
        ws["B5"].font = SECTION_FONT

        metrics = [
            ("Average Book Price", result.summary.average_price),
            ("Median Book Price", result.summary.median_price),
            ("Minimum Book Price", result.summary.min_price),
            ("Maximum Book Price", result.summary.max_price),
        ]

        ws["B7"] = "Metric"
        ws["C7"] = "Value"
        ws["B7"].font = HEADER_FONT
        ws["C7"].font = HEADER_FONT
        ws["B7"].fill = HEADER_FILL
        ws["C7"].fill = HEADER_FILL
        ws["B7"].border = BORDER_HEADER
        ws["C7"].border = BORDER_HEADER

        r = 8
        for lbl, val in metrics:
            c1 = ws.cell(row=r, column=2, value=lbl)
            c2 = ws.cell(row=r, column=3, value=val if val is not None else "N/A")
            c1.font = LABEL_FONT
            c2.font = DATA_FONT
            c1.border = BORDER_ALL
            c2.border = BORDER_ALL
            c2.alignment = Alignment(horizontal="right")
            if val is not None:
                c2.number_format = CURRENCY_FORMAT
            r += 1

        # Top 10 most expensive
        ws["B14"] = "Top 10 Most Expensive Books"
        ws["B14"].font = SECTION_FONT

        top_10 = self.analytics.get_top_expensive(df, n=10)
        self._write_rank_table_detailed(ws, start_row=16, start_col=2, items=top_10)

        # Bottom 10 least expensive
        ws["H14"] = "Top 10 Least Expensive Books"
        ws["H14"].font = SECTION_FONT

        bottom_10 = self.analytics.get_bottom_expensive(df, n=10)
        self._write_rank_table_detailed(ws, start_row=16, start_col=8, items=bottom_10)

        # Excel Bar Chart for Top 10 expensive if records exist
        if top_10:
            chart = BarChart()
            chart.type = "col"
            chart.style = 10
            chart.title = "Top 10 Most Expensive Books (GBP)"
            chart.y_axis.title = "Price (£)"
            chart.x_axis.title = "Book"
            chart.width = 16
            chart.height = 7.5
            chart.legend = None

            data_ref = Reference(ws, min_col=3, min_row=16, max_row=16 + len(top_10))
            cats_ref = Reference(ws, min_col=2, min_row=17, max_row=16 + len(top_10))
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(cats_ref)

            ws.add_chart(chart, "E5")

        self._auto_fit_columns(ws, max_len_cap=40)

    def _write_rank_table_detailed(self, ws, start_row: int, start_col: int, items: list[dict]) -> None:
        """Write rank table with Title, Price, Rating, Availability."""
        headers = ["Title", "Price (GBP)", "Rating", "Status"]
        for c_idx, h in enumerate(headers):
            cell = ws.cell(row=start_row, column=start_col + c_idx, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = BORDER_HEADER
            cell.alignment = Alignment(horizontal="right" if "Price" in h or h == "Rating" else "left")

        curr_row = start_row + 1
        if not items:
            ws.cell(row=curr_row, column=start_col, value="(No data available)").font = Font(name=FONT_FAMILY, italic=True)
            return

        for item in items:
            c1 = ws.cell(row=curr_row, column=start_col, value=item.get("title", "N/A"))
            c2 = ws.cell(row=curr_row, column=start_col + 1, value=item.get("price"))
            c3 = ws.cell(row=curr_row, column=start_col + 2, value=item.get("rating"))
            c4 = ws.cell(row=curr_row, column=start_col + 3, value=item.get("availability", "N/A"))

            for c in (c1, c2, c3, c4):
                c.font = DATA_FONT
                c.border = BORDER_ALL
                if curr_row % 2 == 1:
                    c.fill = ZEBRA_FILL

            c2.number_format = CURRENCY_FORMAT
            c2.alignment = Alignment(horizontal="right")
            c3.alignment = Alignment(horizontal="center")
            curr_row += 1

    # -------------------------------------------------------------------------
    # Sheet 4: Rating Analysis
    # -------------------------------------------------------------------------
    def _build_rating_analysis_sheet(self, wb: openpyxl.Workbook, df, result) -> None:
        """Build rating breakdown and distribution frequency with chart."""
        ws = wb.create_sheet(title="Rating Analysis")
        ws.views.sheetView[0].showGridLines = True

        ws["B2"] = "Customer & Review Rating Analysis"
        ws["B2"].font = TITLE_FONT
        ws["B3"] = f"Average Score: {result.summary.average_rating or 'N/A'} out of 5 stars"
        ws["B3"].font = SUBTITLE_FONT

        ws["B5"] = "Rating Distribution Breakdown"
        ws["B5"].font = SECTION_FONT

        ws["B7"] = "Star Rating"
        ws["C7"] = "Book Count"
        ws["D7"] = "Share of Total"

        for c_idx, col_name in enumerate(["B", "C", "D"]):
            cell = ws[f"{col_name}7"]
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = BORDER_HEADER
            cell.alignment = Alignment(horizontal="right" if col_name in ["C", "D"] else "left")

        dist = result.rating_distribution
        total_books = result.summary.total_books or 1
        r = 8

        # 1 to 5 stars
        for stars in range(1, 6):
            count = dist.get(stars, 0)
            share = (count / total_books) if total_books > 0 and dist else 0.0

            c1 = ws.cell(row=r, column=2, value=f"{stars} Star")
            c2 = ws.cell(row=r, column=3, value=count)
            c3 = ws.cell(row=r, column=4, value=share)

            for c in (c1, c2, c3):
                c.font = DATA_FONT
                c.border = BORDER_ALL
                if r % 2 == 1:
                    c.fill = ZEBRA_FILL

            c2.alignment = Alignment(horizontal="right")
            c2.number_format = INTEGER_FORMAT
            c3.alignment = Alignment(horizontal="right")
            c3.number_format = PERCENT_FORMAT
            r += 1

        # Add Bar Chart if data exists
        if dist:
            chart = BarChart()
            chart.type = "col"
            chart.style = 11
            chart.title = "Books by Star Rating (1 to 5)"
            chart.y_axis.title = "Number of Books"
            chart.x_axis.title = "Rating"
            chart.width = 13
            chart.height = 7
            chart.legend = None

            data_ref = Reference(ws, min_col=3, min_row=7, max_row=12)
            cats_ref = Reference(ws, min_col=2, min_row=8, max_row=12)
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(cats_ref)

            ws.add_chart(chart, "F5")

        self._auto_fit_columns(ws)

    # -------------------------------------------------------------------------
    # Sheet 5: Availability Analysis
    # -------------------------------------------------------------------------
    def _build_availability_sheet(self, wb: openpyxl.Workbook, df, result) -> None:
        """Build stock status and availability distribution table."""
        ws = wb.create_sheet(title="Availability")
        ws.views.sheetView[0].showGridLines = True

        ws["B2"] = "Stock & Inventory Availability Status"
        ws["B2"].font = TITLE_FONT
        ws["B3"] = "Inventory breakdown across the catalog"
        ws["B3"].font = SUBTITLE_FONT

        ws["B5"] = "Inventory Categorization"
        ws["B5"].font = SECTION_FONT

        ws["B7"] = "Status Category"
        ws["C7"] = "Item Count"
        ws["D7"] = "Percentage"

        for col_name in ["B", "C", "D"]:
            cell = ws[f"{col_name}7"]
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = BORDER_HEADER
            cell.alignment = Alignment(horizontal="right" if col_name in ["C", "D"] else "left")

        avail_dist = result.availability_distribution
        total_items = result.summary.total_books or 1
        r = 8

        if avail_dist:
            for status, count in avail_dist.items():
                pct = count / total_items

                c1 = ws.cell(row=r, column=2, value=status)
                c2 = ws.cell(row=r, column=3, value=count)
                c3 = ws.cell(row=r, column=4, value=pct)

                for c in (c1, c2, c3):
                    c.font = DATA_FONT
                    c.border = BORDER_ALL
                    if r % 2 == 1:
                        c.fill = ZEBRA_FILL

                c2.alignment = Alignment(horizontal="right")
                c2.number_format = INTEGER_FORMAT
                c3.alignment = Alignment(horizontal="right")
                c3.number_format = PERCENT_FORMAT
                r += 1
        else:
            c1 = ws.cell(row=r, column=2, value="(No inventory data available)")
            c1.font = Font(name=FONT_FAMILY, italic=True)

        self._auto_fit_columns(ws)


def run_excel_cli() -> None:
    """CLI runner for Excel report generation."""
    parser = argparse.ArgumentParser(description="Generate client-ready Excel analytics report from SQLite.")
    parser.add_argument(
        "--db-path",
        default="data/processed/books.db",
        help="Path to SQLite database (default: data/processed/books.db)",
    )
    parser.add_argument(
        "--output",
        default="reports/exports/books_analytics_report.xlsx",
        help="Target output Excel file path (default: reports/exports/books_analytics_report.xlsx)",
    )

    args = parser.parse_args()
    logging.basicConfig(level=logging.WARNING, format="%(asctime)s [%(levelname)s] %(name)s: %(message)s")

    storage = SQLiteStorage(db_path=args.db_path)
    analytics = BookAnalytics(storage=storage)
    report = ExcelReport(analytics=analytics)

    try:
        output_path = report.generate(output_path=args.output)
        print("Excel report generated successfully.")
        print(f"Output: {output_path}")
    except Exception as exc:
        print(f"Error generating Excel report: {exc}")
        raise SystemExit(1) from exc


if __name__ == "__main__":
    from pathlib import Path
    import sys

    project_root = Path(__file__).resolve().parent.parent.parent
    if str(project_root) not in sys.path:
        sys.path.insert(0, str(project_root))

    run_excel_cli()
